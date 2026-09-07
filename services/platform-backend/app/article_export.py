"""Offline export of trusted, owner-checked article render HTML; never fetch remote URLs."""

from __future__ import annotations

import html
import io
import re
from pathlib import Path
from typing import Any

from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from markdownify import markdownify
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Table, TableStyle

from app.errors import ApiError

MIME = {
    "md": "text/markdown; charset=utf-8",
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}


def styles(tag: Tag) -> dict[str, str]:
    return dict(
        part.strip().split(":", 1) for part in str(tag.get("style", "")).split(";") if ":" in part
    )


def number(value: str, default: float) -> float:
    match = re.match(r"^([0-9.]+)", value.strip())
    return float(match[1]) if match else default


def blocks(root: Tag) -> list[Tag]:
    result = []
    for child in root.children:
        if not isinstance(child, Tag):
            continue
        if child.name in {"section", "div", "ul", "ol", "blockquote"}:
            result.extend(blocks(child))
        else:
            result.append(child)
    return result


def export_article(render_html: str, title: str, format: str) -> bytes:
    if format not in MIME:
        raise ApiError(422, "EXPORT_FORMAT_INVALID", "不支持的下载格式。")
    if len(render_html) > 200_000:
        raise ApiError(413, "EXPORT_TOO_LARGE", "文章内容过大，无法导出。")
    soup = BeautifulSoup(render_html, "html.parser")
    if not soup.find("h1"):
        heading = soup.new_tag("h1")
        heading.string = title
        soup.insert(0, heading)
    if format == "md":
        return markdownify(str(soup), heading_style="ATX").encode("utf-8")
    # Do not silently omit images or allow a converter to fetch arbitrary/private URLs.
    if soup.find("img"):
        raise ApiError(
            422,
            "EXPORT_IMAGES_UNSUPPORTED",
            "当前 PDF/Word 导出暂不支持含图片的文章，请选择 Markdown。",
        )
    output = io.BytesIO()
    if format == "pdf":
        pdfmetrics.registerFont(
            TTFont("STSong-Light", str(Path(__file__).parent / "fonts" / "NotoSansSC.ttf"))
        )
        pdfmetrics.registerFontFamily(
            "STSong-Light",
            normal="STSong-Light",
            bold="NotoSansSC-Bold",
            italic="STSong-Light",
            boldItalic="NotoSansSC-Bold",
        )
        pdfmetrics.registerFont(
            TTFont("NotoSansSC-Bold", str(Path(__file__).parent / "fonts" / "NotoSansSC-Bold.ttf"))
        )
        story: list[Any] = []

        def paragraph(tag: Tag) -> Paragraph:
            css = styles(tag)
            size = number(css.get("font-size", ""), 24 if tag.name == "h1" else 16) * 0.75
            style = ParagraphStyle(
                "article",
                fontName="STSong-Light",
                fontSize=size,
                leading=size * number(css.get("line-height", ""), 1.75),
                textColor=colors.toColor(css.get("color", "#222222")),
                alignment={"left": 0, "center": 1, "right": 2, "justify": 4}.get(
                    css.get("text-align", ""), 0
                ),
                spaceAfter=number(css.get("margin-bottom", ""), 12) * 0.75,
                wordWrap="CJK",
                keepWithNext=tag.name in {"h1", "h2", "h3"},
            )

            def inline(node: Any) -> str:
                if isinstance(node, NavigableString):
                    return html.escape(str(node)).replace("\n", "<br/>")
                if not isinstance(node, Tag):
                    return ""
                text = "".join(inline(child) for child in node.children)
                weight = styles(node).get("font-weight", "").strip().lower()
                name = {"strong": "b", "b": "b", "em": "i", "i": "i", "u": "u", "s": "strike"}.get(
                    node.name
                )
                if weight in {"bold", "bolder"} or number(weight, 400) >= 600:
                    name = "b"
                if name:
                    return f"<{name}>{text}</{name}>"
                if node.name == "br":
                    return "<br/>"
                return text

            text = inline(tag)
            if (
                tag.name in {"h1", "h2", "h3", "th"}
                or number(css.get("font-weight", ""), 400) >= 600
            ):
                text = f"<b>{text}</b>"
            if tag.name == "li":
                index = len(tag.find_previous_siblings("li")) + 1
                text = (f"{index}. " if tag.parent and tag.parent.name == "ol" else "• ") + text
            return Paragraph(text or "<br/>", style)

        for block in blocks(soup):
            if block.name == "table":
                rows = [
                    [paragraph(cell) for cell in row.find_all(["td", "th"], recursive=False)]
                    for row in block.find_all("tr")
                ]
                if rows:
                    table = Table(
                        rows,
                        colWidths=[468 / max(map(len, rows))] * max(map(len, rows)),
                        repeatRows=1,
                    )
                    table.setStyle(
                        TableStyle(
                            [
                                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#d9d9d9")),
                                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                                ("LEFTPADDING", (0, 0), (-1, -1), 8),
                                ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                            ]
                        )
                    )
                    story.append(table)
                    table.spaceAfter = 12
            else:
                story.append(paragraph(block))
        SimpleDocTemplate(
            output,
            pagesize=(612, 792),
            leftMargin=72,
            rightMargin=72,
            topMargin=54,
            bottomMargin=54,
            title=title,
        ).build(story)
    else:
        doc = Document()
        normal = doc.styles["Normal"]
        normal.font.name = "Microsoft YaHei"
        normal.font.size = Pt(12)
        normal.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

        def word_paragraph(tag: Tag, container: Any) -> None:
            css = styles(tag)
            list_style = "List Number" if tag.parent and tag.parent.name == "ol" else "List Bullet"
            p = container.add_paragraph(style=list_style if tag.name == "li" else None)
            p.alignment = {"left": 0, "center": 1, "right": 2, "justify": 3}.get(
                css.get("text-align", ""), 0
            )
            p.paragraph_format.space_after = Pt(number(css.get("margin-bottom", ""), 12) * 0.75)
            p.paragraph_format.line_spacing = number(css.get("line-height", ""), 1.75)
            p.paragraph_format.keep_with_next = tag.name in {"h1", "h2", "h3"}
            size = number(css.get("font-size", ""), 24 if tag.name == "h1" else 16) * 0.75
            for text in tag.descendants:
                if isinstance(text, NavigableString):
                    run = p.add_run(str(text))
                    run.font.size = Pt(size)
                    run.bold = tag.name in {"h1", "h2", "h3", "th"} or any(
                        parent.name in {"b", "strong"} for parent in text.parents
                    )
                    run.italic = any(parent.name in {"i", "em"} for parent in text.parents)
                    run.underline = any(parent.name == "u" for parent in text.parents)
                    color = css.get("color", "#222222").strip()
                    if re.fullmatch(r"#[0-9a-fA-F]{6}", color):
                        run.font.color.rgb = RGBColor.from_string(color[1:])
            background = css.get("background-color", css.get("background", "")).strip()
            if re.fullmatch(r"#[0-9a-fA-F]{6}", background):
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), background[1:])
                p._p.get_or_add_pPr().append(shading)

        for block in blocks(soup):
            if block.name == "table":
                rows = [row.find_all(["td", "th"], recursive=False) for row in block.find_all("tr")]
                if rows:
                    table = doc.add_table(rows=0, cols=max(map(len, rows)))
                    table.style = "Table Grid"
                    for cells in rows:
                        row = table.add_row()
                        for index, cell in enumerate(cells):
                            word_paragraph(cell, row.cells[index])
                            empty = row.cells[index].paragraphs[0]._p
                            empty.getparent().remove(empty)
            else:
                word_paragraph(block, doc)
        doc.save(output)
    return output.getvalue()
